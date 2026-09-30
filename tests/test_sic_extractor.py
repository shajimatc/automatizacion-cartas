"""Ejecutar: python -B -m unittest discover -s tests -p test_sic_extractor.py -v.

Pruebas reales: SIC_TEST_DIR permite indicar la carpeta de los dos DOCX.
Por defecto usa ~/Downloads. No se copian documentos corporativos al repositorio.
Si faltan, se informa explícitamente que estas pruebas se omiten.
"""
import os
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document
from app.services.sic_extractor import extract_sic_providers, SICExtractionError


def table(document, headers, rows):
    result = document.add_table(rows=1, cols=len(headers))
    for cell, value in zip(result.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            cell.text = value
    return result


def extract(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return extract_sic_providers(stream)


def section():
    document = Document()
    document.add_paragraph("SELECCIÓN DE PROVEEDORES")
    document.add_paragraph("TABLA Nº8 LISTA DE PROVEEDORES PROPUESTA")
    return document


class StructureTests(unittest.TestCase):
    def test_reordered_columns_normalization_and_exact_name(self):
        doc = Document()
        doc.add_paragraph("  7. selección   de proveedores : ")
        doc.add_paragraph("Tabla Nº99: lista de proveedores propuesta")
        name = '  Empresa "Original" S.R.L.\nSucursal  '
        table(doc, [' razón  social ', 'Nº', ' nit '], [[name, '1', '001234']])
        self.assertEqual(extract(doc), [{'nit': '001234', 'company_name': name}])

    def test_history_and_later_sections_not_used(self):
        doc = Document()
        doc.add_paragraph('ANTECEDENTES')
        table(doc, ['NIT', 'RAZÓN SOCIAL'], [['999', 'Anterior']])
        doc.add_paragraph('SELECCIÓN DE PROVEEDORES')
        table(doc, ['CONCEPTO', 'DESCRIPCIÓN'], [['Filtros', 'Nacional']])
        doc.add_paragraph('Lista de proveedores propuesta: 2 participaron anteriormente.')
        table(doc, ['NIT', 'RAZÓN SOCIAL'], [['123', 'Actual']])
        doc.add_paragraph('BANDA DE PRECIOS')
        table(doc, ['NIT', 'RAZÓN SOCIAL'], [['888', 'Otra sección']])
        self.assertEqual(extract(doc), [{'nit': '123', 'company_name': 'Actual'}])

    def test_multiple_candidates(self):
        doc = section()
        for nit in ['1', '2']:
            table(doc, ['NIT', 'RAZÓN SOCIAL'], [[nit, 'Empresa']])
        with self.assertRaisesRegex(SICExtractionError, '2 listas candidatas'):
            extract(doc)

    def test_incomplete_rows(self):
        for values in [['', 'Empresa'], ['123', '']]:
            with self.subTest(values=values):
                doc = section(); table(doc, ['NIT', 'RAZÓN SOCIAL'], [values])
                with self.assertRaisesRegex(SICExtractionError, 'incompleto'):
                    extract(doc)

    def test_insufficient_headers(self):
        for headers in [['Código', 'Empresa'], ['NIT', 'Empresa']]:
            with self.subTest(headers=headers):
                doc = section(); table(doc, headers, [['123', 'Empresa']])
                with self.assertRaisesRegex(SICExtractionError, 'encabezados'):
                    extract(doc)

    def test_missing_section_or_list(self):
        for title in ['ANTECEDENTES', 'SELECCIÓN DE PROVEEDORES']:
            doc = Document(); doc.add_paragraph(title)
            table(doc, ['NIT', 'RAZÓN SOCIAL'], [['123', 'Empresa']])
            with self.assertRaises(SICExtractionError): extract(doc)

    def test_unlimited_and_repeated_header(self):
        doc = section()
        values = [[str(i), f'Empresa {i}'] for i in range(100)]
        table(doc, ['NIT', 'RAZÓN SOCIAL'], values[:50] + [['NIT', 'RAZÓN SOCIAL']] + values[50:])
        self.assertEqual(len(extract(doc)), 100)

    def test_merged_title_and_name_columns(self):
        doc = section(); t = doc.add_table(rows=3, cols=4)
        t.cell(0, 0).merge(t.cell(0, 3)).text = 'LISTA DE PROVEEDORES PROPUESTA'
        t.cell(1, 0).text = 'Nº'; t.cell(1, 1).text = 'NIT'
        t.cell(1, 2).merge(t.cell(1, 3)).text = 'RAZÓN SOCIAL'
        t.cell(2, 0).text = '1'; t.cell(2, 1).text = '123'
        t.cell(2, 2).merge(t.cell(2, 3)).text = 'Empresa'
        self.assertEqual(extract(doc), [{'nit': '123', 'company_name': 'Empresa'}])

    def test_duplicate_nit_and_invalid_nit(self):
        for rows in [[['123', 'A'], ['123', 'B']], [['12A', 'A']]]:
            doc = section(); table(doc, ['NIT', 'RAZÓN SOCIAL'], rows)
            with self.assertRaises(SICExtractionError): extract(doc)

    def test_vertical_merge_is_ambiguous(self):
        doc = section(); t = table(doc, ['NIT', 'RAZÓN SOCIAL'], [['123', 'A'], ['', 'B']])
        t.cell(1, 0).merge(t.cell(2, 0))
        with self.assertRaisesRegex(SICExtractionError, 'combinadas ambiguas'): extract(doc)

    def test_empty_and_dangling_list(self):
        doc = section()
        with self.assertRaisesRegex(SICExtractionError, 'sin una tabla'): extract(doc)
        table(doc, ['NIT', 'RAZÓN SOCIAL'], [])
        with self.assertRaisesRegex(SICExtractionError, 'sin proveedores'): extract(doc)


class RealSICTests(unittest.TestCase):
    def read_sic(self, filename):
        folder = Path(os.environ.get('SIC_TEST_DIR', str(Path.home() / 'Downloads')))
        path = folder / filename
        if not path.is_file():
            self.skipTest(f'SIC real no disponible: {path}; configure SIC_TEST_DIR.')
        return extract_sic_providers(path)

    def test_sic_10003015(self):
        rows = self.read_sic('10003015 Informe de Inicio Final.docx')
        self.assertEqual(len(rows), 25)
        self.assertEqual(rows[0], {'nit': '1028617020', 'company_name': 'BOLINTER LTDA.'})
        by_nit = {r['nit']: r['company_name'] for r in rows}
        self.assertEqual(by_nit['1015407022'], 'EQUIPO PETROLERO S.A.')
        self.assertEqual(by_nit['1015577024'], 'SERVICIOS PETROLEROS, SOCIEDAD DE RESPONSABILIDAD LIMITADA" (SERVIPETROL LTDA.)')
        self.assertEqual(rows[-1], {'nit': '169034025', 'company_name': 'WET CHEMICAL BOLIVIA S.R.L.'})
        self.assertFalse(any('SERPETROL' in r['company_name'] for r in rows))

    def test_sic_10003341(self):
        rows = self.read_sic('SIC_10003341 final.docx')
        self.assertEqual(len(rows), 9)
        self.assertEqual(rows[0], {'nit': '1015449028', 'company_name': 'BAKER HUGHES INTERNATIONAL BRANCHES LLC - LTDA. SUCURSAL BOLIVIA'})
        by_nit = {r['nit']: r['company_name'] for r in rows}
        self.assertEqual(by_nit['492403029'], 'O&OTECH BOLIVIA S.R.L.')
        self.assertEqual(rows[-1], {'nit': '282190027', 'company_name': 'WELLSER S.R.L.'})
        self.assertFalse(any('MARRIOTT' in r['company_name'] for r in rows))


if __name__ == '__main__':
    unittest.main()
