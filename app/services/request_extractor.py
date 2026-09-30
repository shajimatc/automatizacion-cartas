import re

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _has_section_numbering(paragraph, document):
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    style = paragraph.style

    while numbering is None and style is not None:
        properties = style.element.pPr
        numbering = properties.numPr if properties is not None else None
        style = style.base_style

    if numbering is None or numbering.numId is None:
        return False

    number_id = numbering.numId.val
    level = numbering.ilvl.val if numbering.ilvl is not None else 0
    definitions = document.part.numbering_part.element
    abstract_ids = definitions.xpath(
        f'./w:num[@w:numId="{number_id}"]/w:abstractNumId/@w:val'
    )
    if not abstract_ids:
        return False

    formats = definitions.xpath(
        f'./w:num[@w:numId="{number_id}"]'
        f'/w:lvlOverride[@w:ilvl="{level}"]/w:lvl/w:numFmt/@w:val'
    ) or definitions.xpath(
        f'./w:abstractNum[@w:abstractNumId="{abstract_ids[0]}"]'
        f'/w:lvl[@w:ilvl="{level}"]/w:numFmt/@w:val'
    )
    return bool(formats and formats[0] not in ("bullet", "none"))


def extract_request_number(file_path):
    document = Document(file_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    match = re.search(
        r"Solicitud\s*N[°º]?\s*(\d+)",
        full_text,
        re.IGNORECASE,
    )

    if match:
        return match.group(1)

    return ""
def extract_process_description(file_path):
    document = Document(file_path)
    lines = [
        " ".join(line.split())
        for element in document.element.body.iter(qn("w:p"))
        for line in Paragraph(element, document).text.splitlines()
        if line.strip()
    ]
    label = re.compile(
        r"^(?:descripci[oó]n|objeto|denominaci[oó]n)\s+"
        r"(?:del\s+(?:proceso|servicio|suministro|contrato)"
        r"|de\s+la\s+contrataci[oó]n)\s*[:\-]\s*(.+)$",
        re.IGNORECASE,
    )

    for line in lines:
        match = label.match(line)
        if match:
            return match.group(1).strip()

    # Etiqueta y valor en celdas contiguas: no tomar texto de otra fila.
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            for index in range(len(cells) - 1):
                heading = " ".join(cells[index].text.split()).rstrip(":-")
                if label.match(heading + ": valor"):
                    value = " ".join(cells[index + 1].text.split())
                    if value and cells[index]._tc is not cells[index + 1]._tc:
                        return value

    # Conservar la redacción original y admitir variantes de contratación.
    introduction = (
        r"(?:para\s+)?(?:la\s+)?"
        r"(?:prestaci[oó]n\s+(?:del?\s+)?servicios?"
        r"(?:\s+especializados?)?\s+de"
        r"|contrataci[oó]n\s+(?:(?:del?\s+)?servicios?\s+de|de)"
        r"|(?:adquisici[oó]n|suministro)\s+de)\s+"
    )
    for index, line in enumerate(lines):
        if not re.search(introduction, line, re.IGNORECASE):
            continue
        # Una descripción puede continuar en los párrafos siguientes;
        # no atravesar títulos numerados ni la sección de anexos.
        candidate = line
        for following in lines[index + 1:]:
            if re.match(r"^(?:\d+[.)]|ANEXOS\b)", following, re.IGNORECASE):
                break
            candidate += " " + following
            if re.search(r"\ben\s+adelante\b", candidate, re.IGNORECASE):
                break
        match = re.search(
            introduction + r"(.+?)\s*[,;.(]\s*en\s+adelante\b",
            candidate,
            re.IGNORECASE,
        )
        if match:
            return match.group(1).strip()

        # Sin la cláusula 'en adelante', aceptar solo un título entre comillas.
        match = re.search(
            introduction + r'["“«]([^"”»]+)["”»]',
            line,
            re.IGNORECASE,
        )
        if match:
            return match.group(1).strip()

    return ""


def extract_request_annexes(file_path):
    document = Document(file_path)
    annexes = []
    in_annexes_section = False

    section_title = re.compile(
        r"\bANEXOS\s+DEL\s+PLIEGO\b",
        re.IGNORECASE,
    )
    numbered_section = re.compile(
        r"^(?:\d+(?:\.\d+)*|[IVXLCDM]+)"
        r"(?:[.)º°:\-]\s*|\s+)\S",
        re.IGNORECASE,
    )
    annex_line = re.compile(
        r"^(?:Anexo|Documento)\b",
        re.IGNORECASE,
    )

    # Incluye párrafos dentro de tablas, en el orden del documento.
    for element in document.element.body.iter(qn("w:p")):
        paragraph = Paragraph(element, document)

        for raw_line in paragraph.text.splitlines():
            line = " ".join(raw_line.split())

            if not line:
                continue

            if not in_annexes_section:
                if section_title.search(line):
                    in_annexes_section = True
                continue

            if annex_line.match(line):
                annexes.append(line)
            elif numbered_section.match(line) or _has_section_numbering(
                paragraph, document
            ):
                return annexes

    return annexes
