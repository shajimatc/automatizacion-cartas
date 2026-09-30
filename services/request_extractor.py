import re

from docx import Document


def extract_request_number(file_path):
    document = Document(file_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    match = re.search(
        r"Solicitud\s*N[°º]?\s*(\d+)",
        full_text,
        re.IGNORECASE,
    )

    if match:
        return match.group(1)

    return ""
def extract_process_description(file_path):
    document = Document(file_path)

    full_text = " ".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    match = re.search(
        r"para la prestación del Servicio de\s+(.+?)"
        r",\s+en adelante denominado simplemente",
        full_text,
        re.IGNORECASE,
    )

    if match:
        return match.group(1).strip()

    return ""